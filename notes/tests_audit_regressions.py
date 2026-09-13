"""Régressions des pages et exports identifiées pendant l'audit du projet."""
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.cache import cache
from django.db import IntegrityError
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from eleves.models import Classe, Ecole, Eleve
from paiements.tests.support import TEST_MIDDLEWARE
from .models import AppreciationMaternelle, ClasseNote, CompositionNote, Evaluation, MatiereNote, NoteMensuelle
from .utils_rangs import calculer_rangs_classe_periode


@override_settings(MIDDLEWARE=TEST_MIDDLEWARE)
class PagesNotesAuditTests(TestCase):
    def setUp(self):
        cache.clear()
        self.ecole = Ecole.objects.create(nom="Ecole audit", adresse="Conakry", telephone="+224600001001", directeur="Direction", etat="VALIDE")
        self.autre_ecole = Ecole.objects.create(nom="Autre ecole audit", adresse="Conakry", telephone="+224600001002", directeur="Direction", etat="VALIDE")
        self.user = User.objects.create_user("notes-audit", password="secret")
        profil = self.user.profil
        profil.ecole = self.ecole
        profil.role = "ADMIN"
        profil.is_validated = True
        profil.save()
        self.client.force_login(self.user)
        self.classe = Classe.objects.create(ecole=self.ecole, nom="CM2", niveau="PRIMAIRE_6", annee_scolaire="2026-2027")
        self.classe_note = ClasseNote.objects.create(ecole=self.ecole, nom="CM2", niveau="PRIMAIRE_6", niveau_enseignement="PRIMAIRE", annee_scolaire="2026-2027")
        self.eleve = Eleve.objects.create(matricule="AUD-001", nom="Diallo", prenom="Aminata", sexe="F", classe=self.classe, statut="ACTIF")
        self.matiere = MatiereNote.objects.create(classe=self.classe_note, nom="Mathématiques", code="MATH", coefficient=1)
        NoteMensuelle.objects.create(eleve=self.eleve, matiere=self.matiere, mois="OCTOBRE", note=0, annee_scolaire="2026-2027")
        self.classe_externe = Classe.objects.create(ecole=self.autre_ecole, nom="CM2", niveau="PRIMAIRE_6", annee_scolaire="2026-2027")
        self.notes_externes = ClasseNote.objects.create(ecole=self.autre_ecole, nom="CM2", niveau="PRIMAIRE_6", niveau_enseignement="PRIMAIRE", annee_scolaire="2026-2027")
        self.eleve_externe = Eleve.objects.create(matricule="EXT-AUD-001", nom="Autre", prenom="Eleve", sexe="M", classe=self.classe_externe, statut="ACTIF")

    def test_saisie_guineenne_affiche_la_classe_et_la_note_zero(self):
        response = self.client.get(reverse("notes:saisie_notes_guineen"), {"classe_id": self.classe_note.pk, "eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["eleve_selectionne"], self.eleve)
        self.assertEqual(response.context["notes_mensuelles"]["OCTOBRE"], 0)

    def test_saisie_guineenne_exclut_un_eleve_exterieur(self):
        response = self.client.get(reverse("notes:saisie_notes_guineen"), {"classe_id": self.classe_note.pk, "eleve_id": self.eleve_externe.pk})
        self.assertEqual(response.status_code, 404)

    def test_bulletin_intelligent_affiche_sans_erreur(self):
        response = self.client.get(reverse("notes:bulletin_intelligent", args=[self.eleve.pk, self.classe_note.pk, "OCTOBRE"]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.eleve.prenom)

    def test_bulletin_refuse_une_classe_notes_exterieure(self):
        for route in ("bulletin_intelligent", "bulletin_intelligent_pdf", "bulletin_intelligent_excel"):
            with self.subTest(route=route):
                response = self.client.get(reverse("notes:" + route, args=[self.eleve.pk, self.notes_externes.pk, "OCTOBRE"]))
                self.assertEqual(response.status_code, 404)

    def test_export_complet_pdf_est_un_pdf(self):
        response = self.client.get(reverse("notes:exporter_notes_complet_pdf"), {"classe_id": self.classe_note.pk, "periode": "OCTOBRE"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_exports_complets_refusent_une_autre_ecole(self):
        for route in ("exporter_notes_complet_pdf", "exporter_notes_complet_excel"):
            with self.subTest(route=route):
                response = self.client.get(reverse("notes:" + route), {"classe_id": self.notes_externes.pk})
                self.assertEqual(response.status_code, 404)

    def test_export_excel_conserve_la_moyenne_zero(self):
        response = self.client.get(reverse("notes:exporter_notes_complet_excel"), {"classe_id": self.classe_note.pk, "periode": "OCTOBRE"})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        row = next(row for row in workbook.active.iter_rows(values_only=True) if self.eleve.matricule in row)
        self.assertEqual(float(row[5]), 0)

    def test_rangs_utilisent_la_classe_de_lecole_sans_identifiant_fixe(self):
        # Les anciens identifiants 61 -> 56 venaient d'une autre installation.
        notes = ClasseNote.objects.create(pk=61, ecole=self.ecole, nom="CM1", niveau="PRIMAIRE_5", niveau_enseignement="PRIMAIRE", annee_scolaire="2026-2027")
        classe = Classe.objects.create(ecole=self.ecole, nom="CM1", niveau="PRIMAIRE_5", annee_scolaire="2026-2027")
        Classe.objects.create(pk=56, ecole=self.autre_ecole, nom="Autre niveau", niveau="PRIMAIRE_5", annee_scolaire="2026-2027")
        eleve = Eleve.objects.create(matricule="AUD-CM1", nom="Barry", prenom="Fatou", sexe="F", classe=classe, statut="ACTIF")
        matiere = MatiereNote.objects.create(classe=notes, nom="Français", code="FR", coefficient=1)
        NoteMensuelle.objects.create(eleve=eleve, matiere=matiere, mois="OCTOBRE", note=8, annee_scolaire="2026-2027")
        rangs = calculer_rangs_classe_periode(notes, "OCTOBRE", use_cache=False)
        self.assertEqual(set(rangs), {eleve.pk})
        self.assertEqual(rangs[eleve.pk]["moyenne"], Decimal("8"))


    def _enregistrer_notes(self, **extra):
        data = {"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk, "annee_scolaire": "2026-2027"}
        data.update(extra)
        return self.client.post(reverse("notes:sauvegarder_notes_guineen"), data, content_type="application/json")

    def test_note_invalide_annule_toute_la_saisie(self):
        response = self._enregistrer_notes(notes_mois={"OCTOBRE": {"note": 8}, "NOVEMBRE": {"note": 99}})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve, mois="OCTOBRE").note, 0)
        self.assertEqual(NoteMensuelle.objects.filter(eleve=self.eleve).count(), 1)

    def test_notes_malformees_sont_refusees_sans_ecriture(self):
        for notes in ({"OCTOBRE": {"note": "NaN"}}, {"INCONNU": {"note": 5}}, [5], {"OCTOBRE": {"note": "abc"}}):
            with self.subTest(notes=notes):
                response = self._enregistrer_notes(notes_mois=notes)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve, mois="OCTOBRE").note, 0)
                self.assertEqual(NoteMensuelle.objects.filter(eleve=self.eleve).count(), 1)

    def test_saisie_recalcule_un_classement_deja_en_cache(self):
        self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], 0)
        response = self._enregistrer_notes(notes_mois={"OCTOBRE": {"note": 8}})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], 8)

    def test_compte_sans_ecole_ne_peut_pas_enregistrer_de_notes(self):
        profil = self.user.profil
        profil.ecole = None
        profil.save()
        response = self._enregistrer_notes(notes_mois={"OCTOBRE": {"note": 8}})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve, mois="OCTOBRE").note, 0)

    def test_suppression_preserve_les_notes_des_autres_annees(self):
        ancienne = NoteMensuelle.objects.create(eleve=self.eleve, matiere=self.matiere, mois="OCTOBRE", note=5, annee_scolaire="2025-2026")
        response = self.client.post(reverse("notes:supprimer_notes"), {"matiere_id": self.matiere.pk, "periode": "OCTOBRE"}, content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(NoteMensuelle.objects.filter(pk=ancienne.pk).exists())
        self.assertFalse(NoteMensuelle.objects.filter(matiere=self.matiere, annee_scolaire="2026-2027").exists())

    def test_suppression_recalcule_un_classement_deja_en_cache(self):
        note = NoteMensuelle.objects.get(eleve=self.eleve, mois="OCTOBRE")
        note.note = 8
        note.save()
        self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], 8)
        response = self.client.post(reverse("notes:supprimer_notes"), {"matiere_id": self.matiere.pk, "periode": "OCTOBRE"}, content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], 0)

    def _sauvegarder_lot(self, **extra):
        data = {"matiere_id": self.matiere.pk, "periode": "OCTOBRE",
                "notes": [{"eleve_id": self.eleve.pk, "note": 0}]}
        data.update(extra)
        return self.client.post(reverse("notes:sauvegarder_notes"), data, content_type="application/json")

    def test_saisie_en_lot_retourne_la_note_zero(self):
        response = self._sauvegarder_lot()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["notes_details"][0]["note"], 0)

    def test_saisie_en_lot_refuse_un_eleve_exterieur_sans_ecriture(self):
        response = self._sauvegarder_lot(notes=[{"eleve_id": self.eleve.pk, "note": 8},
                                              {"eleve_id": self.eleve_externe.pk, "note": 8}])
        self.assertEqual(response.status_code, 403)
        self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve).note, 0)
        self.assertFalse(NoteMensuelle.objects.filter(eleve=self.eleve_externe).exists())

    def test_saisie_en_lot_refuse_une_evaluation_dune_autre_matiere(self):
        matiere = MatiereNote.objects.create(classe=self.notes_externes, nom="Math", code="M")
        evaluation = Evaluation.objects.create(matiere=matiere, titre="Autre", type_evaluation="DEVOIR",
                                               periode="OCTOBRE", date_evaluation="2026-10-01")
        response = self._sauvegarder_lot(evaluation_id=evaluation.pk)
        self.assertEqual(response.status_code, 404)
        self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve).note, 0)

    def test_saisie_en_lot_refuse_les_formats_invalides(self):
        for payload in ([1], {"notes": "invalide"}, {"notes": [None]}):
            with self.subTest(payload=payload):
                data = {"matiere_id": self.matiere.pk, "periode": "OCTOBRE"}
                if isinstance(payload, dict):
                    data.update(payload)
                else:
                    data = payload
                response = self.client.post(reverse("notes:sauvegarder_notes"), data, content_type="application/json")
                self.assertEqual(response.status_code, 400)

    def _appreciations(self, data):
        return self.client.post(reverse("notes:sauvegarder_appreciations_maternelle"), data,
                                content_type="application/json")

    def test_appreciations_refusent_un_eleve_exterieur_dans_les_deux_formats(self):
        for payload in (
            {"eleve_id": self.eleve_externe.pk, "matiere_id": self.matiere.pk,
             "appreciations": {"trimestre1": {"appreciation": "A"}}},
            {"appreciations": [{"eleve_id": self.eleve_externe.pk, "matiere_id": self.matiere.pk,
                                "trimestre": "TRIMESTRE_1", "appreciation": "A"}]},
        ):
            with self.subTest(payload=payload):
                self.assertEqual(self._appreciations(payload).status_code, 403)
                self.assertFalse(AppreciationMaternelle.objects.exists())

    def test_appreciation_invalide_annule_toute_la_saisie(self):
        response = self._appreciations({"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk,
            "appreciations": {"trimestre1": {"appreciation": "A"}, "trimestre2": {"appreciation": "INCONNU"}}})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(AppreciationMaternelle.objects.exists())

    def test_appreciations_enregistrent_absence_et_modification(self):
        for valeur, absent in (("A", False), ("B", False), ("B", True)):
            response = self._appreciations({"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk,
                "appreciations": {"trimestre1": {"appreciation": valeur, "absent": absent}}})
            self.assertEqual(response.status_code, 200)
            appreciation = AppreciationMaternelle.objects.get(eleve=self.eleve, matiere=self.matiere)
            self.assertEqual(appreciation.appreciation, "" if absent else valeur)
            self.assertEqual(appreciation.absent, absent)

    def test_compte_sans_ecole_ne_peut_pas_saisir_en_lot_ou_appreciation(self):
        profil = self.user.profil
        profil.ecole = None
        profil.save()
        self.assertEqual(self._sauvegarder_lot().status_code, 403)
        response = self._appreciations({"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk,
                                      "appreciations": {"trimestre1": {"appreciation": "A"}}})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(AppreciationMaternelle.objects.exists())


    def test_saisies_refusent_un_eleve_dune_autre_classe_de_la_meme_ecole(self):
        autre = Classe.objects.create(ecole=self.ecole, nom="CM1", niveau="PRIMAIRE_5", annee_scolaire="2026-2027")
        self.eleve.classe = autre
        self.eleve.save()
        reponses = [
            self._enregistrer_notes(notes_mois={"OCTOBRE": {"note": 8}}),
            self._sauvegarder_lot(),
            self._appreciations({"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk,
                                 "appreciations": {"trimestre1": {"appreciation": "A"}}}),
        ]
        for response in reponses:
            self.assertEqual(response.status_code, 400)
        self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve).note, 0)
        self.assertFalse(AppreciationMaternelle.objects.exists())
        self.assertFalse(Evaluation.objects.exists())

    def test_note_invalide_en_lot_ne_modifie_rien_et_ne_cree_pas_devaluation(self):
        for valeur in (-1, 11, "NaN", "Infinity", "abc", True):
            with self.subTest(valeur=valeur):
                response = self._sauvegarder_lot(notes=[
                    {"eleve_id": self.eleve.pk, "note": 8},
                    {"eleve_id": self.eleve.pk, "note": valeur},
                ])
                self.assertEqual(response.status_code, 400)
                self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve).note, 0)
                self.assertFalse(Evaluation.objects.exists())

    def test_lot_recalcule_les_rangs_et_enregistre_une_absence(self):
        self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], 0)
        response = self._sauvegarder_lot(notes=[{"eleve_id": self.eleve.pk, "note": "8,5"}])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], Decimal("8.5"))
        response = self._sauvegarder_lot(notes=[{"eleve_id": self.eleve.pk, "note": "", "absent": True}])
        self.assertEqual(response.status_code, 200)
        note = NoteMensuelle.objects.get(eleve=self.eleve)
        self.assertTrue(note.absent)
        self.assertEqual(note.note, 0)
        self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], 0)

    def test_lot_composition_secondaire_accepte_20_et_conserve_les_notes_mensuelles(self):
        self.classe_note.niveau_enseignement = "SECONDAIRE"
        self.classe_note.save()
        response = self._sauvegarder_lot(periode="SEMESTRE_1", notes=[{"eleve_id": self.eleve.pk, "note": 20}])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(CompositionNote.objects.get(eleve=self.eleve, periode="SEMESTRE_1").note, 20)
        self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve).note, 0)
        self.assertEqual(response.json()["notes_details"][0]["note"], 20)

    def test_lot_refuse_une_evaluation_dune_autre_periode(self):
        evaluation = Evaluation.objects.create(matiere=self.matiere, titre="Novembre", type_evaluation="DEVOIR",
                                               periode="NOVEMBRE", date_evaluation="2026-11-01")
        response = self._sauvegarder_lot(evaluation_id=evaluation.pk)
        self.assertEqual(response.status_code, 404)

    def test_appreciations_refusent_formats_periodes_et_annees_invalides(self):
        for extra in (
            {"appreciations": [None]}, {"appreciations": "A"}, {"appreciations": {"trimestre4": {"appreciation": "A"}}},
            {"appreciations": {"trimestre1": {"absent": "false"}}}, {"annee_scolaire": "2025-2026"},
            {"eleve_id": "-1"}, {"matiere_id": [1]},
        ):
            with self.subTest(extra=extra):
                data = {"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk,
                        "appreciations": {"trimestre1": {"appreciation": "A"}}}
                data.update(extra)
                self.assertEqual(self._appreciations(data).status_code, 400)
                self.assertFalse(AppreciationMaternelle.objects.exists())

    def test_appreciations_en_lot_refusees_ne_modifient_pas_les_donnees_existantes(self):
        appreciation = AppreciationMaternelle.objects.create(
            eleve=self.eleve, matiere=self.matiere, annee_scolaire="2026-2027",
            trimestre="TRIMESTRE_1", appreciation="B")
        response = self._appreciations({"appreciations": [
            {"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk, "trimestre": "TRIMESTRE_1", "appreciation": "A"},
            {"eleve_id": self.eleve_externe.pk, "matiere_id": self.matiere.pk, "trimestre": "TRIMESTRE_1", "appreciation": "A"},
        ]})
        self.assertEqual(response.status_code, 403)
        appreciation.refresh_from_db()
        self.assertEqual(appreciation.appreciation, "B")
        self.assertEqual(AppreciationMaternelle.objects.count(), 1)

    def test_appreciations_recalculent_les_rangs_apres_modification_et_absence(self):
        self.classe.nom = self.classe_note.nom = "Petite section"
        self.classe.niveau = self.classe_note.niveau = "MATERNELLE"
        self.classe.save()
        self.classe_note.niveau_enseignement = "MATERNELLE"
        self.classe_note.save()
        for valeur, absent, moyenne in (("A", False, 95), ("B", False, 70), ("B", True, 0)):
            response = self._appreciations({"appreciations": [
                {"eleve_id": self.eleve.pk, "matiere_id": self.matiere.pk,
                 "trimestre": "TRIMESTRE_1", "appreciation": valeur, "absent": absent},
            ]})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "TRIMESTRE_1")[self.eleve.pk]["moyenne"], moyenne)
            self.assertEqual(calculer_rangs_classe_periode(self.classe_note, "OCTOBRE")[self.eleve.pk]["moyenne"], moyenne)

    def test_lot_appreciations_et_absence_utilise_le_bon_trimestre(self):
        for valeur, absent in (("A", False), ("A", True)):
            response = self._sauvegarder_lot(periode="TRIMESTRE_2", notes=[
                {"eleve_id": self.eleve.pk, "appreciation": valeur, "absent": absent},
            ])
            self.assertEqual(response.status_code, 200)
            appreciation = AppreciationMaternelle.objects.get(eleve=self.eleve, trimestre="TRIMESTRE_2")
            self.assertEqual(appreciation.appreciation, "" if absent else "A")
            self.assertEqual(appreciation.absent, absent)
        self.assertFalse(Evaluation.objects.exists())

    def test_erreur_decriture_annule_les_notes_et_levaluation(self):
        second = Eleve.objects.create(matricule="AUD-002", nom="Barry", prenom="Moussa",
                                      sexe="M", classe=self.classe, statut="ACTIF")
        original_save = NoteMensuelle.save

        def save_ou_echec(instance, *args, **kwargs):
            if instance.eleve_id == second.pk:
                raise IntegrityError("Échec simulé de la deuxième écriture")
            return original_save(instance, *args, **kwargs)

        with patch.object(NoteMensuelle, "save", new=save_ou_echec):
            response = self._sauvegarder_lot(notes=[
                {"eleve_id": self.eleve.pk, "note": 8}, {"eleve_id": second.pk, "note": 6},
            ])
        self.assertEqual(response.status_code, 500)
        self.assertEqual(NoteMensuelle.objects.get(eleve=self.eleve).note, 0)
        self.assertFalse(NoteMensuelle.objects.filter(eleve=second).exists())
        self.assertFalse(Evaluation.objects.exists())
