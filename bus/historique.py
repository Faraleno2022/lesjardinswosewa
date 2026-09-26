"""Historique des abonnements d'un élève, renouvellement et documents.

Un élève peut cumuler plusieurs abonnements bus et cantine. Ce module
regroupe ces abonnements, prépare le renouvellement à partir du dernier
abonnement (pour ne pas ressaisir les informations de l'enfant) et produit
la liste PDF / Excel ainsi que le carnet d'abonnement.
"""

import calendar
import io
from datetime import timedelta
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

from .models import AbonnementBus, AbonnementCantine

TYPES_ABONNEMENT = {
    'bus': AbonnementBus,
    'cantine': AbonnementCantine,
}
LIBELLES_TYPES = {'bus': 'Bus', 'cantine': 'Cantine'}

# Informations reprises du dernier abonnement lors d'un renouvellement.
CHAMPS_REPRIS = {
    'bus': (
        'montant', 'periodicite', 'alerte_avant_jours', 'zone', 'itineraire',
        'point_arret', 'contact_parent',
    ),
    'cantine': (
        'montant', 'periodicite', 'type_repas', 'alerte_avant_jours',
        'regime_alimentaire', 'allergies', 'contact_parent',
    ),
}

# Durée de chaque type de paiement : (mois, jours). Les tranches n'ont pas
# de durée fixe : leur date d'expiration reste à saisir.
DUREES = {
    'JOURNALIER': (0, 1),
    'HEBDOMADAIRE': (0, 7),
    'MENSUEL': (1, 0),
    'TRIMESTRIEL': (3, 0),
    'ANNUEL': (12, 0),
}

# Au-delà de ce retard, l'enfant a interrompu l'abonnement : le nouveau
# commence aujourd'hui plutôt qu'au lendemain de l'ancienne expiration.
RETARD_MAX_CONTINUITE = timedelta(days=31)

MOIS = (
    '', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin', 'Juillet',
    'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre',
)
JOURS = ('Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi', 'Dimanche')


def ajouter_mois(jour, nombre):
    mois = jour.month - 1 + nombre
    annee = jour.year + mois // 12
    mois = mois % 12 + 1
    return jour.replace(
        year=annee, month=mois,
        day=min(jour.day, calendar.monthrange(annee, mois)[1]),
    )


def date_expiration_pour(date_debut, periodicite):
    """Dernier jour couvert : 1er septembre en mensuel → 30 septembre."""
    duree = DUREES.get(periodicite)
    if not date_debut or duree is None:
        return None
    mois, jours = duree
    return ajouter_mois(date_debut, mois) + timedelta(days=jours) - timedelta(days=1)


def dernier_abonnement(eleve, type_abonnement):
    modele = TYPES_ABONNEMENT[type_abonnement]
    return (
        modele.objects.filter(eleve=eleve)
        .order_by('-date_expiration', '-created_at')
        .first()
    )


def valeurs_renouvellement(eleve, type_abonnement, aujourd_hui=None):
    """Valeurs initiales d'un nouvel abonnement pour un élève déjà abonné.

    Retourne ``None`` lorsque l'élève n'a encore aucun abonnement de ce type.
    """
    precedent = dernier_abonnement(eleve, type_abonnement)
    if precedent is None:
        return None
    aujourd_hui = aujourd_hui or timezone.localdate()
    valeurs = {
        champ: getattr(precedent, champ) for champ in CHAMPS_REPRIS[type_abonnement]
    }
    suite = precedent.date_expiration + timedelta(days=1)
    date_debut = suite if suite >= aujourd_hui - RETARD_MAX_CONTINUITE else aujourd_hui
    valeurs['date_debut'] = date_debut
    valeurs['date_expiration'] = date_expiration_pour(date_debut, precedent.periodicite)
    valeurs['statut'] = TYPES_ABONNEMENT[type_abonnement].Statut.ACTIF
    valeurs['precedent'] = precedent
    return valeurs


def libelle_mois(date_debut, date_expiration):
    debut = f"{MOIS[date_debut.month]} {date_debut.year}"
    if not date_expiration or (date_expiration.year, date_expiration.month) == (
        date_debut.year, date_debut.month
    ):
        return debut
    fin = f"{MOIS[date_expiration.month]} {date_expiration.year}"
    return f"{debut} – {fin}"


def _statut(abonnement, aujourd_hui):
    if abonnement.statut == abonnement.Statut.ACTIF and abonnement.date_expiration < aujourd_hui:
        return 'Expiré'
    return abonnement.get_statut_display()


def lignes_abonnements(eleve, type_abonnement=None):
    """Abonnements bus et cantine de l'élève, du plus ancien au plus récent."""
    aujourd_hui = timezone.localdate()
    lignes = []
    for code, modele in TYPES_ABONNEMENT.items():
        if type_abonnement and code != type_abonnement:
            continue
        for abonnement in modele.objects.filter(eleve=eleve):
            if code == 'bus':
                detail = ' / '.join(
                    v for v in (abonnement.itineraire, abonnement.point_arret) if v
                )
            else:
                detail = abonnement.get_type_repas_display()
            lignes.append({
                'type': code,
                'service': LIBELLES_TYPES[code],
                'objet': abonnement,
                'detail': detail,
                'periodicite': abonnement.get_periodicite_display(),
                'mois': libelle_mois(abonnement.date_debut, abonnement.date_expiration),
                'montant': abonnement.montant or Decimal('0'),
                'date_debut': abonnement.date_debut,
                'date_expiration': abonnement.date_expiration,
                'jour_expiration': JOURS[abonnement.date_expiration.weekday()],
                'jours_restants': max(0, (abonnement.date_expiration - aujourd_hui).days),
                'reference': abonnement.reference_externe,
                'statut': _statut(abonnement, aujourd_hui),
            })
    lignes.sort(key=lambda l: (l['date_debut'], l['service']))
    return lignes


def total_montants(lignes):
    return sum((ligne['montant'] for ligne in lignes), Decimal('0'))


# --------------------------------------------------------------------------
# Exports
# --------------------------------------------------------------------------

def _nom_fichier(prefixe, eleve, extension):
    matricule = (eleve.matricule or str(eleve.pk)).replace('/', '-').replace(' ', '_')
    return f"{prefixe}_{matricule}_{timezone.localdate():%Y%m%d}.{extension}"


def _gnf(valeur):
    return f"{Decimal(valeur or 0):,.0f}".replace(',', ' ')


def export_excel(eleve, lignes):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Abonnements"
    classe = eleve.classe.nom if eleve.classe else ''
    ws.append([f"Abonnements de {eleve.prenom} {eleve.nom}"])
    ws.append([f"Matricule : {eleve.matricule or ''}", f"Classe : {classe}"])
    ws.append([])
    entetes = [
        'N°', 'Service', 'Détail', 'Type de paiement', 'Mois', 'Montant payé (GNF)',
        'Date de début', "Date d'expiration", "Jour d'expiration", 'Jours restants',
        'Référence', 'Statut',
    ]
    ws.append(entetes)
    ligne_entete = ws.max_row
    for cellule in ws[ligne_entete]:
        cellule.font = Font(bold=True, color='FFFFFF')
        cellule.fill = PatternFill('solid', fgColor='1F4E79')
    ws['A1'].font = Font(bold=True, size=13)
    for numero, ligne in enumerate(lignes, start=1):
        ws.append([
            numero, ligne['service'], ligne['detail'], ligne['periodicite'],
            ligne['mois'], float(ligne['montant']), ligne['date_debut'],
            ligne['date_expiration'], ligne['jour_expiration'],
            ligne['jours_restants'], ligne['reference'], ligne['statut'],
        ])
        for colonne in (7, 8):
            ws.cell(row=ws.max_row, column=colonne).number_format = 'DD/MM/YYYY'
        ws.cell(row=ws.max_row, column=6).number_format = '#,##0'
    ws.append(['', 'TOTAL', '', '', '', float(total_montants(lignes))])
    ws.cell(row=ws.max_row, column=6).number_format = '#,##0'
    for cellule in ws[ws.max_row]:
        cellule.font = Font(bold=True)
    largeurs = (5, 10, 28, 16, 26, 18, 13, 15, 15, 13, 20, 12)
    for index, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur
    ws.freeze_panes = ws.cell(row=ligne_entete + 1, column=1)

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="{_nom_fichier("abonnements", eleve, "xlsx")}"'
    )
    return response


def export_pdf(eleve, lignes, carnet=False, libelle_type=''):
    """Liste des abonnements (paysage) ou carnet d'abonnement (portrait).

    Le carnet reprend, mois par mois, le montant payé, les dates de début et
    d'expiration et le jour d'expiration, avec une colonne de visa.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from ecole_moderne.branding import get_pdf_palette

    ecole = getattr(eleve.classe, 'ecole', None)
    palette = get_pdf_palette(ecole)
    primaire = palette.get('primary', colors.HexColor('#1F4E79'))
    pagesize = A4 if carnet else landscape(A4)
    buffer = io.BytesIO()
    titre = "CARNET D'ABONNEMENT" if carnet else "ABONNEMENTS DE L'ÉLÈVE"
    if libelle_type:
        titre += f" — {libelle_type.upper()}"
    doc = SimpleDocTemplate(
        buffer, pagesize=pagesize, title=titre,
        leftMargin=1.3 * cm, rightMargin=1.3 * cm,
        topMargin=1.2 * cm, bottomMargin=1.3 * cm,
    )
    styles = getSampleStyleSheet()
    style_ecole = ParagraphStyle(
        'Ecole', parent=styles['Normal'], fontName='Helvetica-Bold',
        fontSize=12, textColor=primaire,
    )
    style_titre = ParagraphStyle(
        'Titre', parent=styles['Title'], fontSize=15, textColor=primaire,
        spaceBefore=6, spaceAfter=8,
    )
    style_texte = ParagraphStyle('Texte', parent=styles['Normal'], fontSize=9, leading=12)
    style_cellule = ParagraphStyle('Cellule', parent=styles['Normal'], fontSize=8, leading=9.5)
    style_entete = ParagraphStyle(
        'Entete', parent=style_cellule, fontName='Helvetica-Bold',
        textColor=colors.white, alignment=1, fontSize=7.5, leading=9,
    )

    elements = []
    if ecole is not None:
        coordonnees = ' — '.join(
            v for v in (getattr(ecole, 'adresse', ''), getattr(ecole, 'telephone', '')) if v
        )
        elements.append(Paragraph(ecole.nom, style_ecole))
        if coordonnees:
            elements.append(Paragraph(coordonnees, style_texte))
    elements.append(Paragraph(titre, style_titre))

    responsable = getattr(eleve, 'responsable_principal', None)
    contact = getattr(responsable, 'telephone', '') if responsable else ''
    if not contact:
        # À défaut de responsable, le contact saisi sur le dernier abonnement.
        contact = next(
            (l['objet'].contact_parent for l in reversed(lignes) if l['objet'].contact_parent),
            '',
        )
    identite = [
        ['Élève :', f"{eleve.prenom} {eleve.nom}", 'Matricule :', eleve.matricule or ''],
        ['Classe :', eleve.classe.nom if eleve.classe else '', 'Contact parent :', contact or ''],
    ]
    largeur_utile = pagesize[0] - 2.6 * cm
    bloc = Table(identite, colWidths=[
        largeur_utile * 0.13, largeur_utile * 0.37, largeur_utile * 0.17, largeur_utile * 0.33,
    ])
    bloc.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9.5),
        ('BOX', (0, 0), (-1, -1), 0.6, primaire),
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F5F8FC')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements += [bloc, Spacer(1, 0.4 * cm)]

    if carnet:
        entetes = ['N°', 'Mois', 'Service', 'Montant payé', 'Date de début',
                   "Date d'expiration", "Jour d'expiration", 'Référence', 'Visa']
        largeurs = [0.8, 3.6, 1.7, 2.3, 2.1, 2.2, 2.1, 2.2, 1.6]
    else:
        entetes = ['N°', 'Service', 'Détail', 'Type de paiement', 'Mois', 'Montant payé',
                   'Date de début', "Date d'expiration", "Jour d'expiration",
                   'Jours rest.', 'Référence', 'Statut']
        largeurs = [0.8, 1.6, 4.2, 2.2, 4.0, 2.3, 2.0, 2.2, 2.1, 1.5, 2.5, 1.7]

    donnees = [[Paragraph(entete, style_entete) for entete in entetes]]
    for numero, ligne in enumerate(lignes, start=1):
        if carnet:
            donnees.append([
                numero, Paragraph(ligne['mois'], style_cellule), ligne['service'],
                _gnf(ligne['montant']), f"{ligne['date_debut']:%d/%m/%Y}",
                f"{ligne['date_expiration']:%d/%m/%Y}", ligne['jour_expiration'],
                Paragraph(ligne['reference'] or '', style_cellule), '',
            ])
        else:
            donnees.append([
                numero, ligne['service'], Paragraph(ligne['detail'] or '', style_cellule),
                ligne['periodicite'], Paragraph(ligne['mois'], style_cellule),
                _gnf(ligne['montant']), f"{ligne['date_debut']:%d/%m/%Y}",
                f"{ligne['date_expiration']:%d/%m/%Y}", ligne['jour_expiration'],
                ligne['jours_restants'], Paragraph(ligne['reference'] or '', style_cellule),
                ligne['statut'],
            ])
    colonne_montant = 3 if carnet else 5
    total = [''] * len(entetes)
    total[0] = 'TOTAL PAYÉ'
    total[colonne_montant] = _gnf(total_montants(lignes))
    donnees.append(total)

    tableau = Table(
        donnees, colWidths=[l * cm for l in largeurs], repeatRows=1,
        rowHeights=[None] + [0.85 * cm if carnet else None] * (len(donnees) - 1),
    )
    tableau.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), primaire),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.grey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('ALIGN', (colonne_montant, 1), (colonne_montant, -1), 'RIGHT'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F7F7F7')]),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('SPAN', (0, -1), (colonne_montant - 1, -1)),
        ('ALIGN', (0, -1), (0, -1), 'RIGHT'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E8EEF5')),
    ]))
    elements.append(tableau)
    if not lignes:
        elements.append(Spacer(1, 0.3 * cm))
        elements.append(Paragraph("Aucun abonnement enregistré pour cet élève.", style_texte))

    elements.append(Spacer(1, 0.8 * cm))
    signatures = Table(
        [[f"Édité le {timezone.localdate():%d/%m/%Y}", 'Le Comptable', 'Le Parent']],
        colWidths=[largeur_utile / 3] * 3,
    )
    signatures.setStyle(TableStyle([
        ('FONTNAME', (1, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    elements.append(signatures)
    doc.build(elements)

    prefixe = 'carnet_abonnement' if carnet else 'abonnements'
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="{_nom_fichier(prefixe, eleve, "pdf")}"'
    )
    return response
