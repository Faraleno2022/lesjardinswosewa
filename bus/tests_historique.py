from datetime import date, timedelta
from io import BytesIO
from unittest import mock

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from eleves.models import Classe, Ecole, Eleve
from salaires.tests import TEST_MIDDLEWARE

from . import historique
from .models import AbonnementBus, AbonnementCantine

AUJOURD_HUI = date(2026, 10, 5)


def texte_pdf(response):
    return ''.join(p.extract_text() for p in PdfReader(BytesIO(response.content)).pages)


@override_settings(MIDDLEWARE=TEST_MIDDLEWARE)
@mock.patch('bus.historique.timezone.localdate', return_value=AUJOURD_HUI)
class AbonnementsEleveTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.ecole = Ecole.objects.create(
            nom="École test", adresse="Conakry", telephone="+224622000001",
            directeur="Direction",
        )
        cls.autre_ecole = Ecole.objects.create(
            nom="Autre école", adresse="Siguiri", telephone="+224622000002",
            directeur="Direction",
        )
        cls.classe = Classe.objects.create(
            ecole=cls.ecole, nom="CM1", niveau="PRIMAIRE_5", annee_scolaire="2026-2027",
        )
        cls.eleve = Eleve.objects.create(
            matricule="MAT-010", prenom="Aminata", nom="Camara", sexe="F",
            classe=cls.classe,
        )
        classe_autre = Classe.objects.create(
            ecole=cls.autre_ecole, nom="CM1", niveau="PRIMAIRE_5",
            annee_scolaire="2026-2027",
        )
        cls.eleve_autre = Eleve.objects.create(
            matricule="AUT-001", prenom="Ibrahima", nom="Barry", sexe="M",
            classe=classe_autre,
        )
        cls.user = User.objects.create_user("caisse", password="secret")
        cls.user.profil.ecole = cls.ecole
        cls.user.profil.save(update_fields=["ecole"])

        cls.bus_septembre = AbonnementBus.objects.create(
            eleve=cls.eleve, montant=150000, reference_externe="REC-9",
            periodicite=AbonnementBus.Periodicite.MENSUEL,
            date_debut=date(2026, 9, 1), date_expiration=date(2026, 9, 30),
            zone="Kaloum", itineraire="Ligne A", point_arret="Marché",
            contact_parent="+224620000000",
        )
        cls.cantine = AbonnementCantine.objects.create(
            eleve=cls.eleve, montant=200000, periodicite="MENSUEL",
            type_repas=AbonnementCantine.TypeRepas.REPAS_10H,
            date_debut=date(2026, 9, 1), date_expiration=date(2026, 9, 30),
            regime_alimentaire="Sans porc", allergies="Arachides",
            contact_parent="+224620000000",
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_renouvellement_reprend_les_informations_du_bus(self, _):
        valeurs = historique.valeurs_renouvellement(self.eleve, 'bus')
        self.assertEqual(valeurs['precedent'], self.bus_septembre)
        self.assertEqual(valeurs['montant'], 150000)
        self.assertEqual(valeurs['itineraire'], "Ligne A")
        self.assertEqual(valeurs['point_arret'], "Marché")
        # Continuité : octobre commence le lendemain de l'expiration.
        self.assertEqual(valeurs['date_debut'], date(2026, 10, 1))
        self.assertEqual(valeurs['date_expiration'], date(2026, 10, 31))

    def test_renouvellement_apres_longue_interruption_commence_aujourdhui(self, _):
        AbonnementBus.objects.filter(pk=self.bus_septembre.pk).update(
            date_debut=date(2026, 5, 1), date_expiration=date(2026, 5, 31),
        )
        valeurs = historique.valeurs_renouvellement(self.eleve, 'bus')
        self.assertEqual(valeurs['date_debut'], AUJOURD_HUI)

    def test_aucun_renouvellement_sans_abonnement_precedent(self, _):
        self.assertIsNone(historique.valeurs_renouvellement(self.eleve_autre, 'cantine'))

    def test_formulaire_prerempli_pour_un_eleve_deja_abonne(self, _):
        response = self.client.get(
            reverse('bus:creer_abonnement_cantine'), {'eleve': self.eleve.pk},
        )
        self.assertEqual(response.status_code, 200)
        form = response.context['form']
        self.assertEqual(form.initial['regime_alimentaire'], "Sans porc")
        self.assertEqual(form.initial['allergies'], "Arachides")
        self.assertEqual(form.initial['type_repas'], AbonnementCantine.TypeRepas.REPAS_10H)
        self.assertEqual(form.initial['date_debut'], date(2026, 10, 1))
        self.assertContains(response, 'Renouvellement')

    def test_api_eleve_renvoie_le_renouvellement(self, _):
        response = self.client.get(
            reverse('bus:get_eleve_info_json', args=[self.eleve.pk]), {'type': 'bus'},
        )
        data = response.json()
        self.assertEqual(data['nombre_abonnements'], 1)
        self.assertEqual(data['renouvellement']['itineraire'], "Ligne A")
        self.assertEqual(data['renouvellement']['date_debut'], '2026-10-01')

    def test_liste_des_abonnements_de_l_eleve(self, _):
        AbonnementBus.objects.create(
            eleve=self.eleve, montant=150000, periodicite="MENSUEL",
            date_debut=date(2026, 10, 1), date_expiration=date(2026, 10, 31),
        )
        response = self.client.get(reverse('bus:abonnements_eleve', args=[self.eleve.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context['lignes']), 3)
        self.assertEqual(response.context['total'], 500000)
        self.assertContains(response, 'Octobre 2026')
        self.assertContains(response, 'Samedi')  # 31/10/2026

        bus = self.client.get(
            reverse('bus:abonnements_eleve', args=[self.eleve.pk]), {'type': 'bus'},
        )
        self.assertEqual(len(bus.context['lignes']), 2)

    def test_exports_pdf_excel_et_carnet(self, _):
        excel = self.client.get(reverse('bus:abonnements_eleve_excel', args=[self.eleve.pk]))
        self.assertEqual(excel.status_code, 200)
        feuille = load_workbook(BytesIO(excel.content)).active
        valeurs = [c for ligne in feuille.iter_rows(values_only=True) for c in ligne]
        self.assertIn('Septembre 2026', valeurs)
        self.assertIn("Jour d'expiration", valeurs)
        self.assertIn('Mercredi', valeurs)  # 30/09/2026
        self.assertIn(350000.0, valeurs)

        liste = texte_pdf(self.client.get(reverse('bus:abonnements_eleve_pdf', args=[self.eleve.pk])))
        self.assertIn("ABONNEMENTS DE L'ÉLÈVE", liste)
        self.assertIn('MAT-010', liste)

        carnet = texte_pdf(self.client.get(
            reverse('bus:carnet_abonnement_pdf', args=[self.eleve.pk]), {'type': 'cantine'},
        ))
        self.assertIn("CARNET D'ABONNEMENT", carnet)
        self.assertIn('200 000', carnet)
        self.assertIn('30/09/2026', carnet)
        self.assertIn('Mercredi', carnet)
        self.assertNotIn('150 000', carnet)

    def test_eleve_d_une_autre_ecole_interdit(self, _):
        for nom in ('bus:abonnements_eleve', 'bus:abonnements_eleve_excel',
                    'bus:abonnements_eleve_pdf', 'bus:carnet_abonnement_pdf'):
            response = self.client.get(reverse(nom, args=[self.eleve_autre.pk]))
            self.assertEqual(response.status_code, 404, nom)

    def test_type_inconnu(self, _):
        response = self.client.get(
            reverse('bus:abonnements_eleve', args=[self.eleve.pk]), {'type': 'avion'},
        )
        self.assertEqual(response.status_code, 404)


class DateExpirationTests(TestCase):
    def test_dernier_jour_couvert(self):
        self.assertEqual(
            historique.date_expiration_pour(date(2026, 1, 31), 'MENSUEL'), date(2026, 2, 27),
        )
        self.assertEqual(
            historique.date_expiration_pour(date(2026, 9, 1), 'ANNUEL'), date(2027, 8, 31),
        )
        self.assertEqual(
            historique.date_expiration_pour(date(2026, 9, 1), 'HEBDOMADAIRE'),
            date(2026, 9, 1) + timedelta(days=6),
        )
        self.assertIsNone(historique.date_expiration_pour(date(2026, 9, 1), 'T1'))
